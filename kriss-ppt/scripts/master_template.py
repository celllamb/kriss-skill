"""KRISS native PPTX master/layout binding.

HTML retains all presentation artwork and DOM content; only the PPTX exporter
uses slide-master/layout objects for exact-image matches. No master is rasterized
into an ordinary slide shape. This is not a layout redesigning step.
"""
from __future__ import annotations
import hashlib
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU=914400
TEMPLATE = Path(__file__).resolve().parent.parent/'references'/'kriss-wide-master.pptx'
# Indices refer to the supplied original KRISS wide template; asserted by asset
# hash and geometry on every export so a changed PPTX fails closed.
FAMILY_LAYOUT={'cover':0,'standard':10,'divider':6}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _full_bleed(sh, width, height):
    return (sh.shape_type==MSO_SHAPE_TYPE.PICTURE and
            abs(sh.left)<EMU*.02 and abs(sh.top)<EMU*.02 and
            abs(sh.width-width)<EMU*.02 and abs(sh.height-height)<EMU*.02)


class NativeMaster:
    def __init__(self, path:Path=TEMPLATE):
        self.path=Path(path)
        # A Git checkout transports the original binary in lossless shards.
        # Direct exporter use must work even before verify_installation runs.
        if self.path == TEMPLATE and not self.path.exists():
            from restore_reference import restore_master
            restore_master()
        self.sha256=digest(self.path.read_bytes())
        self.prs=Presentation(self.path)
        if len(self.prs.slide_masters)!=1 or len(self.prs.slide_layouts)<14:
            raise ValueError('KRISS template master or layout inventory changed')
        if abs(self.prs.slide_width/EMU-40/3)>.001 or abs(self.prs.slide_height/EMU-7.5)>.001:
            raise ValueError('KRISS native template size must be exactly 16:9')
        self.family={}
        for name,ix in FAMILY_LAYOUT.items():
            layout=self.prs.slide_layouts[ix]
            owner=self.prs.slide_masters[0] if name=='standard' else layout
            full=[s for s in owner.shapes if _full_bleed(s,self.prs.slide_width,self.prs.slide_height)]
            if len(full)!=1:
                raise ValueError(f'{name}: expected exactly one original full-bleed image in inherited master/layout')
            bg=full[0]
            self.family[name]={'layout_index':ix,'asset_sha256':digest(bg.image.blob),
                               'background_shape_id':bg.shape_id,
                               'background_owner':'master' if name=='standard' else 'layout'}
        # Do not ship the original demonstration slides in user output. Keep the
        # native master, layouts, relationships and theme byte-identical in memory.
        ids=self.prs.slides._sldIdLst
        for sldId in list(ids):
            self.prs.part.drop_rel(sldId.rId)
            ids.remove(sldId)

    def add_slide(self, family:str, background_sha256:str):
        if family not in self.family:
            raise ValueError('Unknown KRISS reference slide family: '+str(family))
        m=self.family[family]
        if m['asset_sha256']!=background_sha256:
            raise ValueError(f'{family}: canonical HTML background differs from native PPTX master/layout')
        return self.prs.slides.add_slide(self.prs.slide_layouts[m['layout_index']]),dict(m)

    def identify_inherited_picture(self,slide,element,asset, sx,sy,tolerance_px=2.0):
        """Map by exact bytes AND native geometry. Never infer from image label.

        Unmatched HTML artwork is exported as ordinary editable PPTX image,
        unless it shares bytes with a native layout image at a conflicting
        position, in which case we block rather than silently duplicate it.
        """
        if not asset or not asset.get('sha256'):
            return None
        matches=[]
        for obj in slide.slide_layout.shapes:
            if obj.shape_type!=MSO_SHAPE_TYPE.PICTURE:continue
            if digest(obj.image.blob)!=asset['sha256']:continue
            wanted=element['rect']
            actual={'x':obj.left/EMU/sx,'y':obj.top/EMU/sy,'w':obj.width/EMU/sx,'h':obj.height/EMU/sy}
            difference=max(abs(actual[k]-wanted[k]) for k in ('x','y','w','h'))
            matches.append((difference,obj))
        if not matches:return None
        matches.sort(key=lambda pair:pair[0])
        difference,obj=matches[0]
        if difference>tolerance_px:
            raise ValueError(f'HTML and PPTX layout contain the same image at different geometry: {element["id"]}, maxDelta={difference:.3f}px')
        return {'kind':'master-image','id':element['id'],'shape_id':obj.shape_id,
                'image_sha256':asset['sha256'],'native_source':'layout',
                'rect':element['rect']}

    def verify_output(self,prs):
        if len(prs.slide_masters)!=1 or len(prs.slide_layouts)!=14:
            raise ValueError('KRISS master/layout lost in serialized output')
        return True
