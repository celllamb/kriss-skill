#!/usr/bin/env python3
"""Capture one final DOM state, export semantic text blocks, retain QA evidence.

No PPTX-side padding guesses, width growth, or per-line textbox generation.
Requires Python 3.10+, Playwright, python-pptx and Chromium.
"""
from __future__ import annotations
import argparse, hashlib, json, os, re, shutil, subprocess, sys
from pathlib import Path
from io import BytesIO
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from lxml import etree
import visual_primitives as vp
from paint_preflight import prepare_paint
from reference_fidelity import check_titles, TITLE_JS
from validate_html import validate_file
from source_coverage import check_coverage
from resource_capture import freeze_images, BACKGROUND_CHECK_JS
import importlib.metadata
from asset_prepare import prepare_assets
from marker_capture import capture_markers
from dom_snapshot import capture_dom_paint, bind_paint_order
from reference_chrome import ReferenceChrome
from qa_common import compare_snapshots
from PIL import Image
from package_info import version_metadata
from speaker_notes import write_notes, notes_equal
from master_template import NativeMaster, TEMPLATE

HERE=Path(__file__).resolve().parent
VIEW=(1280,720)
PPT=(40/3,7.5)
IGNORE='.progress,.counter,.toolbar,.hint,.controls,.speaker-notes,[data-speaker-notes],[data-export-ignore="true"]'
CAPTURE_CSS=IGNORE+'{display:none!important}.deck{box-shadow:none!important}'

def sha(p:Path)->str:return hashlib.sha256(p.read_bytes()).hexdigest()
def write_json(p:Path,obj):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
from runtime_prepare import prepare_html

def set_paragraph(p,d,scale):
    p.alignment=vp.pp_align(d.get('textAlign'))
    pp=p._p.get_or_add_pPr();pp.set('marL','0');pp.set('marR','0');pp.set('indent','0');pp.set('fontAlgn','base');pp.set('latinLnBrk','0');pp.set('hangingPunct','0')
    p.space_before=Pt(d.get('spaceBefore',0)*scale)
    p.space_after=Pt(d.get('spaceAfter',0)*scale)
    # Exact CSS line height in physical units; no estimated font-ratio clamping.
    p.line_spacing=Pt(d['lineHeightPx']*scale)
    p.font.name=vp.first_font(d.get('fontFamily'));p.font.size=Pt(d['fontSize']*scale)
    p.font.bold=float(d.get('fontWeight','400'))>=600 if str(d.get('fontWeight','400')).isdigit() else False
    for seg in d['runs']:
        if seg.get('break'):
            p.add_line_break();continue
        r=p.add_run();r.text=seg['text'];vp.set_text_style(r,seg,fs=seg['fontSize']*scale)
        # python-pptx font.name sets a:latin only. Set East Asian/complex-script
        # faces too, otherwise Korean can silently use the theme font.
        rp=r._r.get_or_add_rPr();rp.set('lang',seg.get('lang') or 'en-US');rp.set('kern','0');rp.set('spc',str(round(seg.get('letterSpacingPx',0)*scale*100)))
        for script in ('ea','cs'):
            for old in list(rp):
                if old.tag==qn('a:'+script):rp.remove(old)
            face=OxmlElement('a:'+script);face.set('typeface',vp.first_font(seg.get('fontFamily')));rp.append(face)
        if seg.get('hyperlink'):
            # Preserve the target, not just link-looking text. rPr carries the
            # authored color/underline; do not substitute a theme hyperlink style.
            r.hyperlink.address=seg['hyperlink']
            # OOXML text hyperlinks otherwise use the theme hyperlink color,
            # even with an explicit run fill in some renderers. This Microsoft
            # extension selects the measured run text color (not a new color).
            link = rp.find(qn('a:hlinkClick'))
            extensions = OxmlElement('a:extLst')
            extension = OxmlElement('a:ext')
            extension.set('uri', '{A12FA001-AC4F-418D-AE19-62706E023703}')
            namespace = 'http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor'
            color_policy = etree.Element('{'+namespace+'}hlinkClr', nsmap={'ahyp': namespace})
            color_policy.set('val', 'tx')
            extension.append(color_policy);extensions.append(extension);link.append(extensions)
        if seg.get('verticalAlign') in ('super','sub'):
            r._r.get_or_add_rPr().set('baseline','30000' if seg['verticalAlign']=='super' else '-25000')

def write_text(slide,block,sx,sy,shape=None):
    cb=block['contentRect']
    if shape is None:
        shape=slide.shapes.add_textbox(Inches(cb['x']*sx),Inches(cb['y']*sy),Inches(cb['w']*sx),Inches(cb['h']*sy))
    tf=shape.text_frame;tf.clear();tf.auto_size=MSO_AUTO_SIZE.NONE;tf.word_wrap=block['wrap']!='intrinsic-nowrap'
    tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[block['verticalAnchor']]
    # Offset is derived from final content area relative to actual native path.
    # For a standalone textbox these are zero; a shape includes padding/border once.
    x,y,w,h=[float(v)/914400 for v in (shape.left,shape.top,shape.width,shape.height)]
    # Preset shapes define an intrinsic text rectangle in addition to bodyPr
    # insets. Compensate that *geometric* rectangle, not font-dependent offsets.
    # roundRect text rect inset = radius * 29289 / 100000 (DrawingML guide).
    pg=shape._element.spPr.find(qn('a:prstGeom'))
    preset=pg.get('prst') if pg is not None else None
    ix=iy=0.0
    if preset=='roundRect':
        radius=float(shape.adjustments[0])*min(w,h);ix=iy=radius*29289/100000
    elif preset=='ellipse':
        ix=w*(1-2**-.5)/2;iy=h*(1-2**-.5)/2
    block['presetTextInsetInches']={'x':ix,'y':iy}
    margins={'left':cb['x']*sx-x-ix,'right':x+w-ix-(cb['x']+cb['w'])*sx,
             'top':cb['y']*sy-y-iy,'bottom':y+h-iy-(cb['y']+cb['h'])*sy}
    for k,v in margins.items():
        setattr(tf,'margin_'+k,Inches(v))
    bp=tf._txBody.bodyPr
    # Preserve native text flow and never use automatic width/height enlargement.
    bp.set('wrap','none' if block['wrap']=='intrinsic-nowrap' else 'square');bp.set('anchorCtr','0')
    for k in ('horzOverflow','vertOverflow'):bp.attrib.pop(k,None)
    if block['wrap']=='intrinsic-nowrap':bp.set('horzOverflow','overflow')
    for i,d in enumerate(block['paragraphs']):set_paragraph(tf.paragraphs[0] if i==0 else tf.add_paragraph(),d,sy*72)
    shape.name='kriss:'+block['id']
    return shape

def capture(html:Path,output:Path,chromium:str|None=None)->dict:
    output.parent.mkdir(parents=True,exist_ok=True)
    work=output.with_suffix('.evidence');work.mkdir(exist_ok=True)
    static=validate_file(html)
    if not static['pass']:raise ValueError('Static preflight failed: '+json.dumps(static['errors'],ensure_ascii=False))
    prepared=prepare_html(html.read_text(encoding='utf-8'))
    prepared,asset_records=prepare_assets(prepared,html.parent)
    # Save the exact delivered content state. UI outside the slide is excluded only from capture.
    final_html=html.with_name(html.stem+'.final.html');final_html.write_text(prepared,encoding='utf-8')

    browser_binary=chromium or os.environ.get('CHROMIUM_PATH') or shutil.which('chromium') or shutil.which('chromium-browser')
    master=NativeMaster();prs=master.prs
    manifest={'schema':'kriss-ppt-evidence-v2','native_template_sha256':master.sha256,'native_template_path':str(master.path),'native_template_families':master.family,**version_metadata(),'static_preflight':static,'local_asset_sources':asset_records,'status':'EXPORTED_NOT_RELEASED','input':str(html), 'input_sha256':sha(html),'final_html':str(final_html),'final_html_sha256':sha(final_html),'slides':[], 'toolchain':{}, 'fonts':{}, 'warnings':[]}
    for file in ['master_template.py','master_typography.py','image_identity.py','package_info.py','export-pptx.py','visual_primitives.py','text-block-extract.js','browser-state.js','adaptive-runtime.js','generated-paint-extract.js','design_paint.py','paint_preflight.py','reference_fidelity.py','runtime_prepare.py','source-inventory.js','speaker_notes.py','source_coverage.py','resource_capture.py','validate_html.py','qa_common.py','build_release.py','asset_prepare.py','marker_capture.py','curved_paint.py','dom_snapshot.py','polygon_paint.py','reference_chrome.py']:
        manifest['toolchain'][file]=sha(HERE/file)
    manifest.update({key+'_relative':os.path.relpath(path,work) for key,path in [('input',html),('final_html',final_html),('pptx',output)]})
    manifest['reference_sha256']=sha(HERE.parent/'references/reference.html')
    with sync_playwright() as pw:
        launch={'headless':True,'args':['--no-sandbox','--disable-dev-shm-usage']}
        if browser_binary:launch['executable_path']=browser_binary
        browser=pw.chromium.launch(**launch);manifest['environment']={'chromium':browser.version,'python':sys.version.split()[0]}
        page=browser.new_page(viewport={'width':VIEW[0],'height':VIEW[1]},device_scale_factor=1)
        blocked_urls=[];responses={};request_failures=[];script_errors=[]
        def capture_response(response):
            if response.request.resource_type=='image' and response.ok:
                try:responses[response.url]=response.body()
                except Exception:pass
        page.on('response',capture_response)
        page.on('pageerror',lambda error:script_errors.append(str(error)))
        page.on('requestfailed',lambda request:request_failures.append({'url':request.url,'type':request.resource_type,'failure':request.failure}))
        if os.environ.get('KRISS_OFFLINE')=='1':
            def offline_route(route):
                if route.request.url.startswith(('http://','https://')):
                    blocked_urls.append(route.request.url);route.abort()
                else:route.continue_()
            page.route('**/*',offline_route)
        page.set_content(final_html.read_text(encoding='utf-8'),wait_until='load',timeout=60000);page.add_style_tag(content=CAPTURE_CSS)
        manifest['environment']['packages']={name:importlib.metadata.version(name) for name in ['playwright','python-pptx','Pillow','numpy']}
        manifest['environment']['platform']=sys.platform
        manifest['environment']['asset_origin']='local assets embedded as original bytes; no browser security bypass'
        manifest['environment']['offline']=os.environ.get('KRISS_OFFLINE')=='1'
        captured_titles=[]
        chrome=ReferenceChrome(browser,HERE.parent/'references/reference.html',{'width':VIEW[0],'height':VIEW[1]})
        manifest['reference_chrome']={'pass':True,'slides':[]}
        count=page.locator('.slide').count()
        if not count:raise ValueError('No .slide elements')
        for i in range(count):
            page.evaluate((HERE/'browser-state.js').read_text(),i)
            captured_titles.extend(t for t in page.evaluate(TITLE_JS) if t['slide']==i+1)
            dom_paint=capture_dom_paint(page,i)
            source=page.evaluate((HERE/'source-inventory.js').read_text(),i)
            marker_result=capture_markers(page,i)
            resource_errors=page.evaluate(BACKGROUND_CHECK_JS,i)
            generated=page.evaluate((HERE/'generated-paint-extract.js').read_text(),i)
            text=page.evaluate((HERE/'text-block-extract.js').read_text(),i)
            visual=page.evaluate(vp.EXTRACT_JS,i)
            visual['elements'].extend(generated['elements'])
            visual['svgLines']=source['svgLines'];visual['svgShapes']=source['svgShapes']
            bind_paint_order(text,visual,dom_paint)
            assets,asset_errors=freeze_images(visual,responses,work)
            resource_errors.extend(asset_errors)
            resource_errors.extend(page.evaluate('window.__KRISS_RESOURCE_REPORT||[]'))
            model={'text':text,'visual':visual,'source':source,'dom_paint':dom_paint,'assets':assets,'resource_errors':resource_errors,'adaptive_log':page.evaluate('window.__KRISS_ADAPTIVE_LOG||[]'),'connector_report':page.evaluate('window.__KRISS_CONNECTOR_REPORT||[]')}
            # Save HTML and model from the SAME state, then temporarily hide children for background.
            hp=work/f'html-{i+1:02d}.png';page.locator('.slide').nth(i).screenshot(path=str(hp),animations='disabled')
            chrome_check=chrome.check(page,i,hp.read_bytes())
            manifest['reference_chrome']['slides'].append(chrome_check)
            manifest['reference_chrome']['pass'] &= chrome_check['pass']
            original_background=chrome.original_asset_for(chrome_check)
            if not original_background:
                raise ValueError(f'Slide {i+1}: no verified KRISS reference family background')
            family=chrome_check['family']
            after=page.evaluate((HERE/'text-block-extract.js').read_text(),i)
            bind_paint_order(after,{'elements':[],'svgLines':[],'svgShapes':[]},dom_paint)
            if text!=after:raise RuntimeError(f'Slide {i+1}: DOM changed during snapshot')
            prepare_paint(text,visual)
            model['generated_paint_probes']=generated['probes']
            model['marker_capture']=marker_result
            mp=work/f'model-{i+1:02d}.json'
            page.evaluate('i=>{const s=document.querySelectorAll(".slide")[i];for(const e of s.children){e.__krissVis=[e.style.getPropertyValue("visibility"),e.style.getPropertyPriority("visibility")];e.style.setProperty("visibility","hidden","important");}}',i)
            try:bg=page.locator('.slide').nth(i).screenshot(animations='disabled')
            finally:page.evaluate('i=>{for(const e of document.querySelectorAll(".slide")[i].children){const [v,p]=e.__krissVis;v?e.style.setProperty("visibility",v,p):e.style.removeProperty("visibility");delete e.__krissVis;}}',i)
            s,inherited=master.add_slide(family,original_background['sha256'])
            # HTML artwork is complete, but the PPTX inherits the exact picture
            # from its native master/layout. Never insert slide background bitmap.
            sx,sy=PPT[0]/text['deck']['w'],PPT[1]/text['deck']['h']
            model['native_master']={'family':family,**inherited,'template_sha256':master.sha256}
            merged={b['id']:b for b in text['blocks'] if b['mergeVisual']}
            items=[]
            for e in visual['elements']:
                if 'slide' in e['classes']:continue
                e['kind']='element';items.append(e)
            elements_by_id={e['id']:e for e in visual['elements']}
            for b in text['blocks']:
                if b['id'] not in merged:
                    el=elements_by_id.get(b['id'],{})
                    items.append({'kind':'textblock','block':b,'order':int(b['id'].split('.e')[1].split('.')[0]),'stackPath':b.get('stackPath',el.get('stackPath',[0])),'domPaintOrder':b.get('domPaintOrder')})
            items+=visual['svgLines']+visual['svgShapes'];items.sort(key=vp.zkey)
            objects=[{'id':f's{i+1:02d}.background','kind':'master-background','shape_id':inherited['background_shape_id'],'native_source':inherited['background_owner'],'layout_index':inherited['layout_index'],'asset_sha256':original_background['sha256']}]
            for e in items:
                if e['kind']=='svgshape':
                    sh=vp.add_svg_shape(s,e,sx,sy);objects.append({'id':e['id'],'kind':'svg-shape','shape_id':sh.shape_id});continue
                if e['kind']=='textblock':
                    sh=write_text(s,e['block'],sx,sy);objects.append({'id':e['block']['id'],'kind':'text','shape_id':sh.shape_id});continue
                if e['kind']=='svgline':
                    sh=vp.add_svg_line(s,e,sx,sy);sh.name='kriss:'+e['id'];objects.append({'id':e['id'],'kind':'svg-stroke' if e.get('strokePolygon') else 'connector','shape_id':sh.shape_id});continue
                r=e['rect'];x,y,w,h=r['x']*sx,r['y']*sy,r['w']*sx,r['h']*sy
                sh=None
                if e.get('boxShadow','none')!='none' or e.get('gradient') or e.get('conic') or vp.rgba(e.get('backgroundColor'))[3]>.001 or any(float((e.get('border'+k)or{}).get('width',0))>0 for k in ['Top','Right','Bottom','Left']):
                    sh=vp.make_visual_shape(s,e,sx,sy,html);sh.name='kriss:'+e['id']
                    from polygon_paint import clipped_inset_bands
                    for band in clipped_inset_bands(s,e,sx,sy):objects.append({'id':e['id'],'kind':'shadow-band','shape_id':band.shape_id})
                    for edge in vp.add_side_accents(s,e,sx,sy):objects.append({'id':e['id'],'kind':'border-edge','shape_id':edge.shape_id})
                    objects.append({'id':e['id'],'kind':'shape-text' if e['id'] in merged else 'shape','shape_id':sh.shape_id})
                    if e.get('conic'):
                        from curved_paint import conic_sectors
                        for sector in conic_sectors(s,e,sx,sy):objects.append({'id':e['id'],'kind':'fill-sector','shape_id':sector.shape_id})
                if e['id'] in merged:
                    if sh is None:raise RuntimeError('Missing visual owner '+e['id'])
                    write_text(s,merged[e['id']],sx,sy,shape=sh)
                for src in ([e['backgroundUrl']] if e.get('backgroundUrl') else [])+([e['src']] if e['tag']=='img' and e.get('src') else []):
                    asset=assets.get(src)
                    inherited_picture=master.identify_inherited_picture(s,e,asset,sx,sy)
                    if inherited_picture:
                        inherited_picture['layout_index']=inherited['layout_index']
                        objects.append(inherited_picture)
                        continue
                    ok=vp.add_picture_stream(s,BytesIO((work/asset['path']).read_bytes()),x,y,w,h) if asset else False
                    if not ok:manifest['warnings'].append({'slide':i+1,'code':'missing-image','id':e['id']})
                    else:s.shapes[-1].name='kriss:image:'+e['id'];objects.append({'id':e['id'],'kind':'image','shape_id':s.shapes[-1].shape_id})
            write_notes(s,source)
            model['native_objects']=objects
            model['coverage']=check_coverage(source,objects,text)
            write_json(mp,model)
            manifest['slides'].append({'slide':i+1,'html_snapshot':hp.name,'html_snapshot_sha256':sha(hp),'model':mp.name,'model_sha256':sha(mp),'original_background_sha256':original_background['sha256'] if original_background else None,'background_mode':'native-master-inheritance','objects':objects,'warnings':text['warnings']})
            for font in text['fontFamilies']:
                if font not in manifest['fonts']:
                    if shutil.which('fc-match'):
                        match=subprocess.run(['fc-match','-f','%{family}',font],capture_output=True,text=True,check=True).stdout
                        exact=font.casefold() in [v.strip().casefold() for v in match.split(',')]
                        manifest['fonts'][font]={'resolved':match,'exact_family':exact or font in ('sans-serif','serif','monospace')}
                    else:manifest['fonts'][font]={'resolved':'unverified','exact_family':False}
        manifest['reference_title']=check_titles(browser,page,HERE.parent/'references/reference.html',{'width':VIEW[0],'height':VIEW[1]},os.environ.get('KRISS_OFFLINE')=='1',captured_titles)
        manifest['environment']['blocked_resource_urls']=sorted(set(blocked_urls))
        manifest['request_failures']=request_failures
        manifest['script_errors']=script_errors
        # Re-open the delivered file. Toggle .active as an ordinary presentation
        # does; do NOT call the export adapter or force adaptive layout here.
        replay=browser.new_page(viewport={'width':VIEW[0],'height':VIEW[1]},device_scale_factor=1)
        if os.environ.get('KRISS_OFFLINE')=='1':
            replay.route('**/*',offline_route)
        replay.set_content(final_html.read_text(encoding='utf-8'),wait_until='load',timeout=60000)
        replay.add_style_tag(content=CAPTURE_CSS)
        reopen=[]
        for rec in manifest['slides']:
            n=rec['slide']-1
            replay.evaluate('''async i=>{
              document.querySelectorAll('.slide').forEach((s,n)=>s.classList.toggle('active',n===i));
              await new Promise(r=>requestAnimationFrame(r));
              let timer;
              try { await Promise.race([window.__KRISS_LAYOUT_SETTLED,new Promise((_,reject)=>{timer=setTimeout(()=>reject(new Error('Reopened HTML did not settle within 30 seconds')),30000);})]); }
              finally {clearTimeout(timer);}
              await new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)));
            }''',n)
            actual=replay.evaluate((HERE/'source-inventory.js').read_text(),n)
            model=json.loads((work/rec['model']).read_text())
            original=model['source']
            rp=work/f'reopen-{n+1:02d}.png'
            replay.locator('.slide').nth(n).screenshot(path=str(rp),animations='disabled')
            geometry_equal=actual['nodes']==original['nodes'] and actual['overflow']==original['overflow']
            snapshot_equal=sha(rp)==rec['html_snapshot_sha256']
            styles_equal=actual['styles']==original['styles']
            speaker_notes_equal=notes_equal(actual,original)
            render_parity=compare_snapshots(Image.open(work/rec['html_snapshot']).convert('RGB'),Image.open(rp).convert('RGB'),model)
            reopen.append({'slide':n+1,'source_geometry_equal':geometry_equal,'computed_styles_equal':styles_equal,'speaker_notes_equal':speaker_notes_equal,'snapshot_equal':snapshot_equal,'snapshot':rp.name,'render_parity':render_parity,'pass':geometry_equal and styles_equal and speaker_notes_equal and render_parity['pass']})
        manifest['html_reopen']={'pass':all(x['pass'] for x in reopen),'slides':reopen}
        replay.close();chrome.close();browser.close()
    master.verify_output(prs)
    prs.save(output);manifest['pptx']=str(output);manifest['pptx_sha256']=sha(output)
    write_json(work/'manifest.json',manifest)
    print(f'EXPORTED (release QA still required): {count} slides; {output}')
    return manifest

def main():
    ap=argparse.ArgumentParser();ap.add_argument('input_html',type=Path);ap.add_argument('output_pptx',type=Path,nargs='?');ap.add_argument('--chromium')
    a=ap.parse_args();capture(a.input_html.resolve(),(a.output_pptx or a.input_html.with_suffix('.pptx')).resolve(),a.chromium)
if __name__=='__main__':main()
