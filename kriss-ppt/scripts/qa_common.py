"""Evidence-driven QA. A testable boundary is not a claim of visual perfection."""
from __future__ import annotations
import hashlib,json,math,re,shutil,subprocess,tempfile
from pathlib import Path
from collections import Counter
import numpy as np
from PIL import Image,ImageFilter,ImageChops
from pptx import Presentation
from pptx.oxml.ns import qn
from speaker_notes import check_native_notes, validate_notes_source
from image_identity import check_native_images

EMU=914400

def load(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def save(p,obj):Path(p).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def evidence_for(pptx):return Path(pptx).with_suffix('.evidence')
def attested_path(manifest: dict, key: str, evidence: Path) -> Path | None:
    relative=manifest.get(key+'_relative')
    if relative:
        candidate=(evidence/relative).resolve()
        if candidate.is_file():return candidate
    value=manifest.get(key)
    return Path(value) if value else None


def verify_inputs(pptx: Path, evidence: Path):
    errors=[]
    try:m=load(evidence/'manifest.json')
    except (OSError,ValueError) as error:return {},['Unreadable manifest: '+str(error)]
    if m.get('schema')!='kriss-ppt-evidence-v2':errors.append('Missing current fail-closed evidence schema')
    if not m.get('slides'):errors.append('No completed captured slides in the manifest')
    def bound(path: Path | None, expected: str | None, label: str):
        if not path or not expected or not path.is_file() or sha(path)!=expected:
            errors.append(label+' hash mismatch or file unavailable')
            return False
        return True
    for key in ('input','final_html'):
        bound(attested_path(m,key,evidence),m.get(key+'_sha256'),key)
    bound(Path(pptx),m.get('pptx_sha256'),'PPTX')
    for slide in m.get('slides',[]):
        label='slide '+str(slide.get('slide','?'))
        model_path=evidence/slide.get('model','missing')
        if not bound(model_path,slide.get('model_sha256'),label+' model'):continue
        try:model=load(model_path)
        except (ValueError,OSError) as error:errors.append(label+' unreadable model: '+str(error));continue
        for asset in model.get('assets',{}).values():
            bound(evidence/asset['path'],asset.get('sha256'),label+' captured asset')
        bound(evidence/slide.get('html_snapshot','missing'),slide.get('html_snapshot_sha256'),label+' HTML snapshot')
    return m,errors

def preset_inset(sh):
    pg=sh._element.spPr.find(qn('a:prstGeom'));name=pg.get('prst') if pg is not None else None
    w,h=sh.width/EMU,sh.height/EMU
    if name=='roundRect':
        v=float(sh.adjustments[0])*min(w,h)*29289/100000;return v,v
    if name=='ellipse':return w*(1-2**-.5)/2,h*(1-2**-.5)/2
    return 0,0

def native_qa(pptx:Path,evidence:Path)->dict:
    m,errors=verify_inputs(pptx,evidence)
    if errors:return {'pass':False,'errors':errors,'slides':[]}
    prs=Presentation(pptx);rows=[]
    if len(prs.slides)!=len(m['slides']):errors.append('slide count mismatch')
    for slide,rec in zip(prs.slides,m['slides']):
        model=load(evidence/rec['model']);tm=model['text'];sx=(prs.slide_width/EMU)/tm['deck']['w'];sy=(prs.slide_height/EMU)/tm['deck']['h'];bad=[]
        notes_check=check_native_notes(slide,model.get('source',{}))
        bad.extend(notes_check['errors'])
        image_check=check_native_images(slide,model,sx,sy)
        bad.extend(image_check['errors'])
        names={}
        expected_background=rec.get('original_background_sha256')
        if expected_background:
            backgrounds=[sh for sh in slide.shapes if sh.name=='kriss:background']
            if len(backgrounds)!=1 or not hasattr(backgrounds[0],'image') or hashlib.sha256(backgrounds[0].image.blob).hexdigest()!=expected_background:
                bad.append({'error':'original template image bytes changed inside PPTX'})
        for sh in slide.shapes:names.setdefault(sh.name,[]).append(sh)
        expected={b['id'] for b in tm['blocks']}
        for b in tm['blocks']:
            candidates=[sh for sh in names.get('kriss:'+b['id'],[]) if sh.has_text_frame and sh.text.strip()]
            if len(candidates)!=1:bad.append({'id':b['id'],'error':'semantic text must map to exactly one native object','count':len(candidates)});continue
            sh=candidates[0];tf=sh.text_frame
            if tf.text!=b['plainText']:bad.append({'id':b['id'],'error':'text mismatch','expected':b['plainText'],'actual':tf.text})
            if len(tf.paragraphs)!=len(b['paragraphs']):bad.append({'id':b['id'],'error':'paragraph count mismatch'})
            if b['mergeVisual'] and sh._element.spPr.find(qn('a:prstGeom')) is None and sh._element.spPr.find(qn('a:custGeom')) is None:bad.append({'id':b['id'],'error':'simple visual text box was not merged'})
            ix,iy=preset_inset(sh);actual={'x':(sh.left/EMU+ix+tf.margin_left/EMU)/sx,'y':(sh.top/EMU+iy+tf.margin_top/EMU)/sy,
                'w':(sh.width/EMU-2*ix-tf.margin_left/EMU-tf.margin_right/EMU)/sx,'h':(sh.height/EMU-2*iy-tf.margin_top/EMU-tf.margin_bottom/EMU)/sy}
            delta={k:round(actual[k]-b['contentRect'][k],6) for k in actual}
            if max(abs(v) for v in delta.values())>.10:bad.append({'id':b['id'],'error':'content-box mismatch','delta_px':delta})
            for p,d in zip(tf.paragraphs,b['paragraphs']):
                from visual_primitives import pp_align
                if p.alignment!=pp_align(d.get('textAlign')):bad.append({'id':b['id'],'error':'alignment mismatch'})
                expected_breaks=sum(bool(r.get('break')) for r in d['runs']);actual_breaks=len(p._p.findall(qn('a:br')))
                if expected_breaks!=actual_breaks:bad.append({'id':b['id'],'error':'manual-linebreak mismatch'})
                runs=[x for x in d['runs'] if not x.get('break')]
                if len(runs)!=len(p.runs):bad.append({'id':b['id'],'error':'rich run count mismatch'})
                for r,x in zip(p.runs,runs):
                    if abs(r.font.size.pt-x['fontSize']*sy*72)>.015:bad.append({'id':b['id'],'error':'font size mismatch'})
                    try:
                        actual_link=r.hyperlink.address or None
                    except (KeyError, ValueError) as error:
                        actual_link=None
                        bad.append({'id':b['id'],'error':'broken native hyperlink relationship','detail':str(error)})
                    if x.get('hyperlink'):
                        color_tag='{http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor}hlinkClr'
                        link_color=r._r.find('.//'+color_tag)
                        if link_color is None or link_color.get('val') != 'tx':
                            bad.append({'id':b['id'],'error':'hyperlink text-color policy missing'})
                        from visual_primitives import rgb_color
                        expected_color=rgb_color(x.get('color','#000'))
                        try:
                            if r.font.color.rgb != expected_color:
                                bad.append({'id':b['id'],'error':'hyperlink text color mismatch'})
                        except (AttributeError, TypeError):
                            bad.append({'id':b['id'],'error':'hyperlink text color is not explicit RGB'})
                        if bool(r.font.underline) != ('underline' in str(x.get('textDecoration',''))):
                            bad.append({'id':b['id'],'error':'hyperlink underline mismatch'})
                    if actual_link != (x.get('hyperlink') or None):
                        bad.append({'id':b['id'],'error':'hyperlink target mismatch',
                                    'expected':x.get('hyperlink'),'actual':actual_link})
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text.strip() and sh.name.removeprefix('kriss:') not in expected:bad.append({'shape_id':sh.shape_id,'error':'unmapped native text'})
        source_lines=model['visual']['svgLines']
        for line in source_lines:
            hit=names.get('kriss:'+line['id'],[])
            if line.get('strokePolygon'):
                if len(hit)!=1 or hit[0]._element.tag!=qn('p:sp') or hit[0]._element.spPr.find(qn('a:custGeom')) is None:
                    bad.append({'id':line['id'],'error':'missing editable affine SVG stroke'})
                else:
                    sh=hit[0];pts=sh._element.spPr.findall('.//'+qn('a:pt'))
                    expected=line['strokePolygon'];got=[(sh.left/EMU/sx+int(p.get('x'))/EMU/sx,sh.top/EMU/sy+int(p.get('y'))/EMU/sy) for p in pts]
                    if len(got)!=len(expected) or any(abs(x-y)>.1 for p,q in zip(got,expected) for x,y in zip(p,q)):
                        bad.append({'id':line['id'],'error':'affine SVG stroke vertices mismatch'})
                continue
            if len(hit)!=1 or hit[0]._element.tag!=qn('p:cxnSp'):bad.append({'id':line['id'],'error':'missing native connector'})
            else:
                sh=hit[0];bounds=[min(line['x1'],line['x2']),min(line['y1'],line['y2']),abs(line['x2']-line['x1']),abs(line['y2']-line['y1'])]
                vals=[sh.left/EMU/sx,sh.top/EMU/sy,sh.width/EMU/sx,sh.height/EMU/sy]
                if max(abs(x-y) for x,y in zip(bounds,vals))>.1:bad.append({'id':line['id'],'error':'connector coordinates mismatch'})
        # Native design contract: one-side CSS must never acquire a full outline.
        from design_paint import border_plan, normalized_stops
        from visual_primitives import rgba
        for item in model['visual']['elements']:
            if 'slide' in item.get('classes',[]):continue
            plan=border_plan(item,rgba)
            if item.get('clipPolygon') or item.get('borderWedge'):
                hits=names.get('kriss:'+item['id'],[])
                if len(hits)!=1 or hits[0]._element.spPr.find(qn('a:custGeom')) is None:bad.append({'id':item['id'],'error':'missing native polygon/wedge'})
                continue
            if item.get('conic'):
                for i,_ in enumerate(item['conic']['sectors']):
                    sectors=names.get('kriss:sector:'+item['id']+':'+str(i),[])
                    if len(sectors)!=1 or sectors[0]._element.spPr.find(qn('a:custGeom')) is None:bad.append({'id':item['id'],'error':'missing native conic sector','sector':i})
            if not (plan['edges'] or item.get('gradient')):continue
            hits=[sh for sh in names.get('kriss:'+item['id'],[]) if sh._element.tag==qn('p:sp')]
            if len(hits)!=1:
                bad.append({'id':item['id'],'error':'missing/unbound native paint owner'});continue
            sh=hits[0];sp=sh._element.spPr;line=sp.find(qn('a:ln'))
            if plan['mode']!='uniform' and (line is None or line.find(qn('a:noFill')) is None):
                bad.append({'id':item['id'],'error':'unintended full-perimeter border'})
            spec=plan.get('gradient') if plan['mode']=='single-side-fill' else item.get('gradient')
            if spec:
                grad=sp.find(qn('a:gradFill'))
                if grad is None:bad.append({'id':item['id'],'error':'missing native gradient fill'})
                else:
                    lin=grad.find(qn('a:lin'))
                    if lin is None or int(lin.get('ang'))!=round(spec['angle']*60000)%21600000:
                        bad.append({'id':item['id'],'error':'gradient angle mismatch'})
                    stops=grad.find(qn('a:gsLst'));expected_positions=[round(pos*100000) for pos,c in normalized_stops(spec,rgba)]
                    if [int(gs.get('pos')) for gs in stops]!=expected_positions:
                        bad.append({'id':item['id'],'error':'gradient stops mismatch'})
            if plan['mode']=='edges':
                for side in plan['edges']:
                    edge=names.get('kriss:edge:'+item['id']+':'+side,[])
                    if len(edge)!=1:bad.append({'id':item['id'],'error':'missing independent border edge','side':side})
        # Independent source inventory, not a list derived from generated shapes.
        from source_coverage import check_coverage
        if not model.get('source'):
            bad.append({'error':'missing independent source inventory'})
        else:
            live=[]
            for obj in rec['objects']:
                hits=[sh for sh in slide.shapes if sh.shape_id==obj['shape_id']]
                expected_name=('kriss:background' if obj['kind']=='background' else
                    'kriss:image:'+obj['id'] if obj['kind']=='image' else 'kriss:'+obj['id'])
                if obj['kind'] in ('border-edge','fill-sector','shadow-band'):continue
                if len(hits)==1 and hits[0].name==expected_name:live.append(obj)
            coverage=check_coverage(model['source'],live,tm)
            bad.extend(coverage['errors'])
        for item in model['visual'].get('svgShapes',[]):
            hits=names.get('kriss:'+item['id'],[])
            if len(hits)!=1:
                bad.append({'id':item['id'],'error':'missing native SVG shape'});continue
            sh=hits[0];r=item['rect']
            vals=[sh.left/EMU/sx,sh.top/EMU/sy,sh.width/EMU/sx,sh.height/EMU/sy]
            if max(abs(x-y) for x,y in zip(vals,[r['x'],r['y'],r['w'],r['h']]))>.1:
                bad.append({'id':item['id'],'error':'SVG geometry mismatch'})
            pg=sh._element.spPr.find(qn('a:prstGeom'))
            preset='ellipse' if item['tag'] in ('circle','ellipse') else ('roundRect' if item.get('radius') else 'rect')
            if pg is None or pg.get('prst')!=preset:bad.append({'id':item['id'],'error':'SVG primitive type mismatch'})
        # Order is captured by Chromium independently, not inferred from exporter shape count.
        order=model.get('dom_paint',{}).get('order',{})
        prior=None
        for sh in slide.shapes:
            identity=sh.name.removeprefix('kriss:image:') if sh.name.startswith('kriss:image:') else sh.name.removeprefix('kriss:')
            key=order.get(identity)
            if not key:continue
            current=(int(key['layer']),int(key['sequence']))
            if prior and current<prior[0]:bad.append({'id':identity,'error':'native paint-order inversion','previous':prior[1]})
            prior=(current,identity)
        # Images and visuals are verified by correspondence, not a token text count.
        actual_ids={sh.shape_id for sh in slide.shapes}
        for obj in rec['objects']:
            if obj['shape_id'] not in actual_ids:bad.append({'id':obj['id'],'error':'missing mapped object'})
        rows.append({'slide':rec['slide'],'expected_text_blocks':len(tm['blocks']),'speaker_notes':notes_check,'images':image_check,'native_connectors':sum(not bool(x.get('strokePolygon')) for x in source_lines),'native_svg_strokes':sum(bool(x.get('strokePolygon')) for x in source_lines),'errors':bad,'pass':not bad})
    return {'pass':not errors and all(r['pass'] for r in rows),'errors':errors,'slides':rows}

def html_qa(evidence:Path)->dict:
    m=load(evidence/'manifest.json');rows=[];font_errors=[k for k,v in m['fonts'].items() if not v['exact_family']]
    for rec in m['slides']:
        model=load(evidence/rec['model']);bad=list(model['text']['warnings']);warn=[];
        bad.extend(validate_notes_source(model.get('source',{})))
        if not model.get('source') or not model.get('coverage'):bad.append({'code':'missing-source-coverage'})
        else:
            bad.extend(model['coverage']['errors'])
            for overflow in model['source'].get('overflow',[]):
                if not overflow.get('allowed'):bad.append({'code':'unresolved-adaptive-overflow',**overflow})
        bad.extend(model.get('resource_errors',[]))
        dw,dh=model['text']['deck']['w'],model['text']['deck']['h']
        for b in model['text']['blocks']:
            cb=b['contentRect']
            if cb['w']<=0 or cb['h']<=0:bad.append({'id':b['id'],'error':'empty content box'})
            if cb['x']<-1 or cb['y']<-1 or cb['x']+cb['w']>dw+1 or cb['y']+cb['h']>dh+1:bad.append({'id':b['id'],'error':'text box outside slide'})
            for p in b['paragraphs']:
                if p['fontSize']<16:warn.append({'id':b['id'],'warning':'review small font at presentation scale','font_px':p['fontSize']})
                for r in p['glyphRects']:
                    # Font boxes can extend beyond CSS leading vertically; horizontal
                    # clipping is unambiguous and blocking. Visual parity covers ink.
                    if r['x']<cb['x']-2 or r['x']+r['w']>cb['x']+cb['w']+2:bad.append({'id':b['id'],'error':'HTML text exceeds content width'})
        for c in model['connector_report']:
            if c.get('error'):bad.append(c)
        rows.append({'slide':rec['slide'],'pass':not bad,'errors':bad,'warnings':warn})
    errors=list(m.get('warnings',[]))
    errors.extend({'code':'browser-request-failed',**r} for r in m.get('request_failures',[]))
    errors.extend({'code':'browser-script-error','detail':r} for r in m.get('script_errors',[]))
    if not m.get('static_preflight',{}).get('pass'):errors.append({'code':'static-preflight-not-passed'})
    if not m.get('html_reopen',{}).get('pass'):errors.append({'code':'delivered-html-state-mismatch'})
    if not m.get('reference_chrome',{}).get('pass'):errors.append({'code':'original-reference-chrome-not-preserved'})
    return {'pass':not font_errors and not errors and m.get('reference_title',{'pass':False})['pass'] and all(r['pass'] for r in rows),
            'font_errors':font_errors,'errors':errors,'html_reopen':m.get('html_reopen'),
            'static_preflight':m.get('static_preflight'),'reference_chrome':m.get('reference_chrome'),'reference_title':m.get('reference_title',{'applicable':False}),'slides':rows}

def render_pptx(pptx:Path,destination:Path)->list[Path]:
    lo=shutil.which('libreoffice') or shutil.which('soffice');ppm=shutil.which('pdftoppm')
    if not lo or not ppm:raise RuntimeError('LibreOffice and pdftoppm are required for release parity')
    destination.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='kriss-render-') as tmp:
        t=Path(tmp);profile=(t/'profile').as_uri()
        result=subprocess.run([lo,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(t),str(pptx.resolve())],capture_output=True,text=True,timeout=120)
        pdf=t/(pptx.stem+'.pdf')
        if result.returncode or not pdf.exists():raise RuntimeError('PPTX render failed: '+result.stdout+result.stderr)
        shutil.copy2(pdf,destination/'render.pdf')
        subprocess.run([ppm,'-png','-scale-to-x','1280','-scale-to-y','720',str(pdf),str(t/'slide')],check=True,capture_output=True,timeout=120)
        paths=sorted(t.glob('slide-*.png'),key=lambda p:int(p.stem.split('-')[-1]))
        out=[]
        for i,p in enumerate(paths,1):q=destination/f'pptx-{i:02d}.png';shutil.copy2(p,q);out.append(q)
        return out

def region_compare(a:Image.Image,b:Image.Image,rect:dict,pad:int=5)->dict:
    x0=max(0,math.floor(rect['x'])-pad);y0=max(0,math.floor(rect['y'])-pad);x1=min(a.width,math.ceil(rect['x']+rect['w'])+pad);y1=min(a.height,math.ceil(rect['y']+rect['h'])+pad)
    if x1<=x0 or y1<=y0:return {'pass':False,'reason':'empty-region'}
    aa=np.asarray(a.crop((x0,y0,x1,y1)).convert('RGB')).astype(np.int16);bb=np.asarray(b.crop((x0,y0,x1,y1)).convert('RGB')).astype(np.int16)
    # Compare local foreground with a 2px anti-aliasing tolerance; do not register
    # or translate either image. Complete loss of a small label is still a failure.
    bg=np.median(aa.reshape(-1,3),axis=0);ma=(np.abs(aa-bg).max(2)>55);mb=(np.abs(bb-bg).max(2)>55)
    def dilate(mask):return np.asarray(Image.fromarray((mask*255).astype('uint8')).filter(ImageFilter.MaxFilter(5)))>0
    missing=float((ma&~dilate(mb)).sum()/max(1,ma.sum()));extra=float((mb&~dilate(ma)).sum()/max(1,mb.sum()))
    mean=float(abs(aa-bb).mean())
    return {'missing_fraction':missing,'extra_fraction':extra,'mean':mean,'pass':missing<=.20 and extra<=.20}

def compare_snapshots(aa: Image.Image, bb: Image.Image, model: dict) -> dict:
    """Identical numerical/local gates for native and delivered-HTML renders.

    Hash equality remains provenance, not a substitute for geometry/style and
    global/local pixel checks. Thresholds are deliberately not configurable here.
    """
    if aa.size != bb.size:
        return {'pass':False,'error':'snapshot size mismatch'}
    d=abs(np.asarray(aa).astype(np.int16)-np.asarray(bb).astype(np.int16));mean=float(d.mean());major=float((d.max(2)>30).mean());strong=float((d.max(2)>60).mean())
    regions=[]
    for block in model['text']['blocks']:
        for pi,p in enumerate(block['paragraphs']):
            gr=p['glyphRects']
            if not gr:continue
            x=min(r['x'] for r in gr);y=min(r['y'] for r in gr);rr=max(r['x']+r['w'] for r in gr);bottom=max(r['y']+r['h'] for r in gr)
            v=region_compare(aa,bb,{'x':x,'y':y,'w':rr-x,'h':bottom-y});regions.append({'id':block['id'],'paragraph':pi,**v})
    # Presence checks cover every visible shape/image/SVG/connector, not text only.
    visual_regions=[]
    visual_items=[x for x in model['visual']['elements'] if 'slide' not in x.get('classes',[])] + model['visual'].get('svgShapes',[])
    for item in visual_items:
        visual_regions.append({'id':item['id'],**region_compare(aa,bb,item['rect'])})
    for item in model['visual']['svgLines']:
        half=max(.5,float(item.get('strokeWidth',1))/2)
        box={'x':min(item['x1'],item['x2'])-half,'y':min(item['y1'],item['y2'])-half,
             'w':abs(item['x2']-item['x1'])+2*half,'h':abs(item['y2']-item['y1'])+2*half}
        visual_regions.append({'id':item['id'],**region_compare(aa,bb,box)})
    # A text-only local gate cannot detect a dropped gradient or added outline.
    # Use the existing global numerical limits on each changed paint region;
    # do not register images or increase tolerances for this implementation.
    paint_regions=[]
    for item in model['visual']['elements']:
        if 'slide' in item.get('classes',[]) or not(item.get('gradient') or item.get('conic') or item.get('borderMode')=='single-side-fill'):continue
        r=item['rect'];x=max(0,math.ceil(r['x']));y=max(0,math.ceil(r['y']));right=min(aa.width,math.floor(r['x']+r['w']));bottom=min(aa.height,math.floor(r['y']+r['h']))
        if right<=x or bottom<=y:continue
        da=abs(np.asarray(aa.crop((x,y,right,bottom))).astype(np.int16)-np.asarray(bb.crop((x,y,right,bottom))).astype(np.int16))
        pm=float(da.mean());pj=float((da.max(2)>30).mean());ps=float((da.max(2)>60).mean())
        paint_regions.append({'id':item['id'],'mean':pm,'major':pj,'strong':ps,'pass':pm<=3 and pj<=.02 and ps<=.012})
    row={'mean':mean,'major':major,'strong':strong,'global_pass':mean<=3 and major<=.02 and strong<=.012,'regions':regions}
    row['paint_regions']=paint_regions;row['visual_regions']=visual_regions;row['pass']=row['global_pass'] and all(r['pass'] for r in regions) and all(r['pass'] for r in paint_regions) and all(r['pass'] for r in visual_regions)
    return row

def parity_qa(pptx:Path,evidence:Path)->dict:
    m,errors=verify_inputs(pptx,evidence)
    if errors:return {'pass':False,'errors':errors,'slides':[]}
    paths=render_pptx(pptx,evidence/'rendered');rows=[]
    if len(paths)!=len(m['slides']):errors.append('rendered slide count mismatch')
    montage=Image.new('RGB',(1280,360*len(paths)),'white')
    for rec,pp in zip(m['slides'],paths):
        aa=Image.open(evidence/rec['html_snapshot']).convert('RGB');bb=Image.open(pp).convert('RGB')
        if aa.size!=bb.size:errors.append('snapshot size mismatch');continue
        model=load(evidence/rec['model'])
        row={'slide':rec['slide'],**compare_snapshots(aa,bb,model)};rows.append(row)
        montage.paste(aa.resize((640,360)),(0,(rec['slide']-1)*360));montage.paste(bb.resize((640,360)),(640,(rec['slide']-1)*360))
    montage.save(evidence/'comparison.png')
    return {'pass':not errors and all(r['pass'] for r in rows),'errors':errors,'tolerances':{'mean':3,'major':.02,'strong':.012,'local_missing':.20,'local_extra':.20,'local_dilation_px':2},'slides':rows,'renderer':'LibreOffice; not Microsoft PowerPoint'}

def regression_qa(baseline:Path,current:Path,allowed:set[int])->dict:
    bm=load(baseline/'manifest.json');cm=load(current/'manifest.json');rows=[];errors=[]
    if len(bm['slides'])!=len(cm['slides']):errors.append('slide count changed')
    for old,new in zip(bm['slides'],cm['slides']):
        a=load(baseline/old['model']);b=load(current/new['model'])
        # Compare final semantic scene, excluding log timing/paths.
        changed=a['text']!=b['text'] or a['visual']!=b['visual']
        render_changed=sha(baseline/old['html_snapshot'])!=sha(current/new['html_snapshot'])
        n=new['slide'];ok=not(changed or render_changed) or n in allowed
        rows.append({'slide':n,'model_changed':changed,'html_render_changed':render_changed,'allowed':n in allowed,'pass':ok})
    return {'pass':not errors and all(r['pass'] for r in rows),'errors':errors,'slides':rows}
