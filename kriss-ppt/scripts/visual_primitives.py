#!/usr/bin/env python3
"""Translate final rendered KRISS HTML DOM into editable native PowerPoint objects.

HTML is the only source of truth. The browser completes adaptive layout first; this
exporter reads final DOM geometry/computed style and translates it to native PPTX
objects. It must not reinterpret layout.
"""
from __future__ import annotations

import argparse, base64, math, os, re, urllib.request
from io import BytesIO
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

PPT_W = 40 / 3
PPT_H = 7.5
VIEW_W = 1280
VIEW_H = 720
PX_TO_PT = 72 / 96

# No deck-specific calibration knobs: the exporter uses frozen browser DOM
# geometry and computed styles directly. Any layout problem must be fixed in the
# canonical HTML/adaptive layer, not by slide/pattern magic offsets here.


def parse_args():
    ap = argparse.ArgumentParser()
    ap.add_argument('input_html')
    ap.add_argument('output_pptx', nargs='?')
    ap.add_argument('--chromium', default=os.environ.get('CHROMIUM_PATH', '/usr/bin/chromium'))
    return ap.parse_args()


def css_px(v):
    try:
        return float(re.match(r'[-+\d.]+', str(v or '0')).group())
    except Exception:
        return 0.0


def rgba(s):
    if not s or s in ('transparent','none'):
        return (255, 255, 255, 0.0)
    s = str(s).strip()
    m = re.match(r'#([0-9a-fA-F]{3,8})$', s)
    if m:
        h = m.group(1)
        if len(h) in (3, 4):
            h = ''.join(c*2 for c in h)
        r, g, b = int(h[0:2],16), int(h[2:4],16), int(h[4:6],16)
        a = int(h[6:8],16)/255 if len(h) >= 8 else 1.0
        return (r,g,b,a)
    m = re.match(r'rgba?\(\s*([\d.]+)[, ]+\s*([\d.]+)[, ]+\s*([\d.]+)(?:\s*[,/]\s*([\d.]+))?\s*\)', s, re.I)
    if m:
        return (int(round(float(m.group(1)))), int(round(float(m.group(2)))), int(round(float(m.group(3)))), float(m.group(4) or 1.0))
    return (0,0,0,1.0)


def rgb_color(s):
    r,g,b,_ = rgba(s)
    return RGBColor(max(0,min(255,r)), max(0,min(255,g)), max(0,min(255,b)))


def _set_alpha_on_srgb(parent, alpha: float):
    """Apply DrawingML alpha to the first srgbClr child under parent."""
    try:
        colors = parent.xpath('.//a:srgbClr')
        if not colors:
            return
        color = colors[0]
        for child in list(color):
            if child.tag.endswith('}alpha'):
                color.remove(child)
        if alpha < .999:
            a = OxmlElement('a:alpha')
            a.set('val', str(int(max(0,min(1,alpha))*100000)))
            color.append(a)
    except Exception:
        pass


def set_run_alpha_and_spacing(run, opacity=1.0, letter_spacing_px=0.0):
    try:
        rPr = run._r.get_or_add_rPr()
        _set_alpha_on_srgb(rPr, float(opacity or 1.0))
        # DrawingML character spacing is 1/100 pt. CSS px -> pt = * 72/96.
        spc = int(round(float(letter_spacing_px or 0.0) * PX_TO_PT * 100))
        if spc:
            rPr.set('spc', str(spc))
        elif 'spc' in rPr.attrib:
            del rPr.attrib['spc']
    except Exception:
        pass


def first_font(s):
    """Return the first CSS-declared family without deck-specific substitution.

    Actual font-face matching belongs in a general font resolution layer; this
    exporter must not force KRISS/BIPM-specific replacements based on the current regression deck.
    """
    s = str(s or "Arial").split(",")[0].strip().strip("\"\'")
    return s or "Arial"


def add_alpha_to_solid_fill(fill, alpha: float):
    if alpha >= .999:
        return
    try:
        solid = fill._xPr.solidFill
        color = solid[0]
        for child in list(color):
            if child.tag.endswith('}alpha'):
                color.remove(child)
        a = OxmlElement('a:alpha')
        a.set('val', str(int(max(0,min(1,alpha))*100000)))
        color.append(a)
    except Exception:
        pass


def remove_theme_style(shape):
    """Remove PowerPoint theme style refs that can introduce implicit fills/lines/effects.

    The exporter writes explicit fill, line, and shadow values from the frozen
    browser DOM. Leaving p:style in place lets Office/LibreOffice apply theme
    effect references (notably shadows) that are not present in CSS.
    """
    try:
        el = shape._element
        for child in list(el):
            if child.tag == qn('p:style'):
                el.remove(child)
    except Exception:
        pass


def clear_shape_effects(shape):
    try:
        spPr = shape._element.spPr
        for child in list(spPr):
            if child.tag in (qn('a:effectLst'), qn('a:effectDag')):
                spPr.remove(child)
        spPr.append(OxmlElement('a:effectLst'))
    except Exception:
        pass


def apply_css_shadow(shape, box_shadow: str | None):
    clear_shape_effects(shape)
    if not box_shadow or box_shadow == 'none' or 'inset' in str(box_shadow):
        return
    m = re.search(r'(rgba?\([^)]*\)|#[0-9a-fA-F]{3,8})\s+(-?[\d.]+)px\s+(-?[\d.]+)px\s+([\d.]+)px(?:\s+(-?[\d.]+)px)?', str(box_shadow))
    if not m:
        return
    color_css, ox, oy, blur = m.group(1), float(m.group(2)), float(m.group(3)), float(m.group(4))
    r,g,b,a = rgba(color_css)
    try:
        spPr = shape._element.spPr
        for child in list(spPr):
            if child.tag in (qn('a:effectLst'), qn('a:effectDag')):
                spPr.remove(child)
        effect = OxmlElement('a:effectLst')
        sh = OxmlElement('a:outerShdw')
        emu_per_px = 914400 / 96
        dist = math.hypot(ox, oy) * emu_per_px
        ang = (math.degrees(math.atan2(oy, ox)) % 360.0) * 60000
        sh.set('blurRad', str(int(max(0, blur) * emu_per_px)))
        sh.set('dist', str(int(dist)))
        sh.set('dir', str(int(ang)))
        sh.set('algn', 'ctr')
        sh.set('rotWithShape', '0')
        clr = OxmlElement('a:srgbClr')
        clr.set('val', f'{r:02X}{g:02X}{b:02X}')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', str(int(max(0,min(1,a))*100000)))
        clr.append(alpha); sh.append(clr); effect.append(sh); spPr.append(effect)
    except Exception:
        pass


def set_shape_fill(shape, color_css, opacity=1.0):
    r,g,b,a = rgba(color_css)
    a *= opacity
    if a <= .001:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(r,g,b)
    add_alpha_to_solid_fill(shape.fill, a)


def set_line(line, color_css, width_px, opacity=1.0):
    r,g,b,a = rgba(color_css)
    a *= opacity
    if width_px <= 0 or a <= .001:
        line.fill.background()
        return
    line.color.rgb = RGBColor(r,g,b)
    line.width = Pt(width_px * PX_TO_PT)
    try:
        _set_alpha_on_srgb(line._get_or_add_ln(), a)
    except Exception:
        pass


def image_stream(src: str, base_dir: Path):
    if not src:
        return None
    if src.startswith('data:image/'):
        try:
            return BytesIO(base64.b64decode(src.split(',',1)[1]))
        except Exception:
            return None
    if src.startswith('file:'):
        from urllib.parse import urlparse, unquote
        p = Path(unquote(urlparse(src).path))
        return str(p) if p.exists() else None
    if re.match(r'^https?://', src, re.I):
        try:
            with urllib.request.urlopen(src, timeout=15) as r:
                return BytesIO(r.read())
        except Exception:
            return None
    p = (base_dir / src).resolve()
    return str(p) if p.exists() else None


def add_picture(slide, src, base_dir, x,y,w,h):
    stream = image_stream(src, base_dir)
    if stream is None:
        return False
    try:
        slide.shapes.add_picture(stream, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        return True
    except Exception:
        return False

def add_picture_stream(slide, stream, x, y, w, h):
    try:
        stream.seek(0)
        slide.shapes.add_picture(stream, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        return True
    except Exception:
        return False


def shape_type_for(item, w, h, sx):
    classes = item.get('classes', []) or []
    force_circle = False  # Computed CSS, not class names or metadata, defines the visible shape.
    radius_raw = str(item.get('borderRadius') or '')
    radius_px = css_px(item.get('borderRadiusPx') if item.get('borderRadiusPx') is not None else radius_raw)
    radius_in = radius_px * sx
    # CSS border-radius:50% is an ellipse for a non-square box and a circle for a
    # square box. Preserve that semantic as a native OVAL instead of approximating
    # it with a rounded rectangle. Large absolute radii (e.g. 999px pills) remain
    # rounded rectangles/capsules.
    pct = re.match(r'\s*([\d.]+)%', radius_raw)
    percent_ellipse = bool(pct and float(pct.group(1)) >= 50.0)
    if force_circle or percent_ellipse:
        return MSO_SHAPE.OVAL
    if radius_in > .002:
        return MSO_SHAPE.ROUNDED_RECTANGLE
    return MSO_SHAPE.RECTANGLE


def set_text_style(run, item, fs=None):
    font = run.font
    font.name = first_font(item.get('fontFamily'))
    font.size = Pt(fs if fs is not None else float(item.get('fontSize') or 12) * PX_TO_PT)
    fw = str(item.get('fontWeight','400'))
    try: bold = float(fw) >= 600
    except Exception: bold = fw.lower() == 'bold'
    font.bold = bool(bold)
    font.italic = item.get('fontStyle') in ('italic','oblique')
    font.underline = 'underline' in str(item.get('textDecoration',''))
    font.color.rgb = rgb_color(item.get('color','#000'))
    # Preserve computed opacity. Do not apply class/pattern-specific visual
    # compensation factors such as cover/title/closing alpha multipliers.
    set_run_alpha_and_spacing(run, float(item.get('opacity', 1.0)), item.get('letterSpacingPx', 0.0))


def pp_align(css_align):
    a = str(css_align or '').lower()
    if a in ('center','-webkit-center'):
        return PP_ALIGN.CENTER
    if a in ('right','end'):
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def make_visual_shape(slide, item, sx, sy, html_path, with_text=False):
    """Map native fill/outline without turning a one-sided border into four sides."""
    from design_paint import border_plan, parse_linear_gradient, set_gradient
    plan=border_plan(item,rgba)
    if item.get('clipPolygon') or item.get('borderWedge'):plan={'mode':'polygon','edges':{}}
    r=item['rect'];x,y,w,h=r['x']*sx,r['y']*sy,r['w']*sx,r['h']*sy
    if plan['mode']=='uniform':
        bw=float(plan['outline']['width'])
        x+=bw*sx/2;y+=bw*sy/2;w=max(.001,w-bw*sx);h=max(.001,h-bw*sy)
    from curved_paint import requires_corner_path,set_corner_path
    custom_corners=requires_corner_path(item)
    if item.get('clipPolygon') or item.get('borderWedge'):custom_corners=True
    st=MSO_SHAPE.RECTANGLE if custom_corners else shape_type_for(item,w,h,sx)
    sh=slide.shapes.add_shape(st,Inches(x),Inches(y),Inches(w),Inches(h))
    remove_theme_style(sh)
    if item.get('clipPolygon') or item.get('borderWedge'):
        from polygon_paint import set_polygon_path
        set_polygon_path(sh,item.get('clipPolygon') or item['borderWedge']['points'],r['w'],r['h'],sx,sy)
    elif custom_corners:set_corner_path(sh,item,sx,sy,float(plan['outline']['width'])/2 if plan['mode']=='uniform' else 0.)
    if st==MSO_SHAPE.ROUNDED_RECTANGLE and len(sh.adjustments):
        radius=float(item.get('borderRadiusPx') or 0)
        if plan['mode']=='uniform':radius=max(0,radius-float(plan['outline']['width'])/2)
        sh.adjustments[0]=max(0,min(.5,radius/max(1e-6,min(w/sx,h/sy))))
    opacity=float(item.get('opacity',1))
    gradient=item.get('gradient')
    if not gradient and not item.get('conic') and not item.get('unsupportedPaint') and item.get('backgroundImage','none')!='none' and not item.get('backgroundUrl'):
        gradient=parse_linear_gradient(item['backgroundImage'],r['w'],r['h'])
    if plan['mode']=='single-side-fill':
        set_gradient(sh,plan['gradient'],rgba,opacity)
    elif gradient:
        set_gradient(sh,gradient,rgba,opacity)
    else:
        set_shape_fill(sh,item['borderWedge']['color'] if item.get('borderWedge') else item.get('backgroundColor'),opacity)
    if plan['mode']=='uniform':
        b=plan['outline'];set_line(sh.line,b['color'],float(b['width']),opacity)
        _set_dash(sh.line,b.get('style'))
    else:
        sh.line.fill.background()
    apply_css_shadow(sh,'none' if item.get('flatShadow') else item.get('boxShadow'))
    if abs(float(item.get('rotation') or 0))>.01:sh.rotation=float(item['rotation'])
    return sh


def _set_dash(line,style):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    if style=='dashed':line.dash_style=MSO_LINE_DASH_STYLE.DASH
    elif style=='dotted':line.dash_style=MSO_LINE_DASH_STYLE.ROUND_DOT


def add_side_accents(slide,item,sx,sy):
    """Explicit edges for mixed-side styles. Never invent an enclosing outline.

    The usual single-side card/cell is encoded in the primary native shape fill.
    Mixed borders use at most four editable native edges, named under their owner.
    Unsupported nonuniform rounded corners are blocked by paint preflight.
    """
    from design_paint import border_plan
    plan=border_plan(item,rgba)
    if item.get('borderWedge') or item.get('clipPolygon'):return []
    if plan['mode']!='edges':return []
    from curved_paint import corner_radii
    if any(x>0 or y>0 for x,y in corner_radii(item)):
        from curved_paint import rounded_edges
        return rounded_edges(slide,item,sx,sy,plan['edges'])
    r=item['rect'];out=[]
    for side,b in plan['edges'].items():
        x,y,w,h=r['x'],r['y'],r['w'],r['h'];bw=float(b['width'])
        if side=='top':points=(x,y+bw/2,x+w,y+bw/2)
        elif side=='bottom':points=(x,y+h-bw/2,x+w,y+h-bw/2)
        elif side=='left':points=(x+bw/2,y,x+bw/2,y+h)
        else:points=(x+w-bw/2,y,x+w-bw/2,y+h)
        edge=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(points[0]*sx),Inches(points[1]*sy),Inches(points[2]*sx),Inches(points[3]*sy))
        remove_theme_style(edge);clear_shape_effects(edge)
        set_line(edge.line,b['color'],bw,float(item.get('opacity',1)));_set_dash(edge.line,b.get('style'))
        edge.name='kriss:edge:'+item['id']+':'+side
        out.append(edge)
    return out


def add_svg_line(slide, item, sx, sy):
    if item.get('strokePolygon'):
        from curved_paint import new_path_shape
        points=item['strokePolygon'];x=min(p[0] for p in points);y=min(p[1] for p in points)
        rect={'x':x,'y':y,'w':max(p[0] for p in points)-x,'h':max(p[1] for p in points)-y}
        shape,pen=new_path_shape(slide,rect,sx,sy,item['stroke'],item.get('opacity',1)*item.get('strokeOpacity',1),'kriss:'+item['id'])
        pen.move((points[0][0]-x,points[0][1]-y))
        for px,py in points[1:]:pen.line((px-x,py-y))
        pen.close();return shape
    x1,y1,x2,y2 = item['x1']*sx,item['y1']*sy,item['x2']*sx,item['y2']*sy
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    remove_theme_style(line);clear_shape_effects(line)
    set_line(line.line, item.get('stroke','#000'), float(item.get('strokeWidth',1)), float(item.get('opacity',1))*float(item.get('strokeOpacity',1)))
    line.line._get_or_add_ln().set('cap',{'round':'rnd','square':'sq'}.get(item.get('linecap'),'flat'))
    return line


EXTRACT_JS = r"""
(slideIndex) => {
  const slides=[...document.querySelectorAll('.slide')];
  const slide=slides[slideIndex];
  const deck=document.querySelector('.deck')||slide.parentElement;
  const dr=deck.getBoundingClientRect(), sr=slide.getBoundingClientRect();
  const dom=[slide,...slide.querySelectorAll('*')];
  const idOf=e=>'s'+String(slideIndex+1).padStart(2,'0')+'.e'+String(dom.indexOf(e)).padStart(4,'0');
  const skipSelector='script,style,.progress,.counter,.toolbar,.hint,.controls,.speaker-notes,[data-speaker-notes],[data-export-ignore="true"]';
  const px=s=>Number.parseFloat(s||'0')||0;
  const colorAlpha=s=>{const m=String(s).match(/rgba?\([^,]+,[^,]+,[^,]+(?:,\s*([\d.]+))?\)/i);return m&&m[1]!=null?Number(m[1]):(String(s)==='transparent'?0:1)};
  const classPath=el=>{const out=[];let c=el;while(c&&c!==slide){if(c.classList)out.push(...c.classList);c=c.parentElement}return out};
  const stackPath=el=>{if(el===slide)return [-1000];const vals=[];const chain=[];let cur=el;while(cur&&cur!==slide){chain.push(cur);cur=cur.parentElement}chain.reverse();for(const n of chain){const cs=getComputedStyle(n);const v=cs.zIndex;if(v!=='auto'&&v!==''){vals.push(Number.parseInt(v,10)||0)}}return vals.length?vals:[0]};
  const zi=el=>{const p=stackPath(el);return p[p.length-1]||0};
  const opacityOf=el=>{let v=1;for(let n=el;n&&n!==slide.parentElement;n=n.parentElement)v*=Number(getComputedStyle(n).opacity);return v;};
  const visible=el=>{if(!(el instanceof Element))return false;if(el.matches(skipSelector)||el.closest(skipSelector))return false;const cs=getComputedStyle(el),r=el.getBoundingClientRect();if(cs.display==='none'||cs.visibility==='hidden'||opacityOf(el)<=.001)return false;return r.width>.25&&r.height>.25&&r.right>sr.left&&r.left<sr.right&&r.bottom>sr.top&&r.top<sr.bottom};
  const rect=r=>({x:r.left-sr.left,y:r.top-sr.top,w:r.width,h:r.height});
  const rot=t=>{if(!t||t==='none')return 0;const m=t.match(/^matrix\(([-+\d.e]+),\s*([-+\d.e]+),/i);return m?Math.atan2(Number(m[2]),Number(m[1]))*180/Math.PI:0};
  const bgUrl=v=>{if(!v||v==='none')return null;const m=v.match(/^url\(["']?(.*?)["']?\)$/s);return m?m[1]:null};
  const radiusPx=(v,r)=>{v=String(v||'0');const m=v.match(/([\d.]+)%/);if(m)return Number(m[1])/100*Math.min(r.width,r.height);return px(v)};
  const pad=cs=>({left:px(cs.paddingLeft),right:px(cs.paddingRight),top:px(cs.paddingTop),bottom:px(cs.paddingBottom)});
  const lineHeightPx=cs=>{if(cs.lineHeight==='normal')return px(cs.fontSize)*1.2;return px(cs.lineHeight)};
  const stylePack=el=>{const cs=getComputedStyle(el);return {fontFamily:cs.fontFamily,fontSize:px(cs.fontSize),fontWeight:cs.fontWeight,fontStyle:cs.fontStyle,color:cs.color,textDecoration:cs.textDecorationLine,lineHeightPx:lineHeightPx(cs),letterSpacingPx:px(cs.letterSpacing),opacity:Number(cs.opacity||1),textAlign:cs.textAlign,display:cs.display}};
  const hasVisual=el=>{const cs=getComputedStyle(el);return colorAlpha(cs.backgroundColor)>.001||cs.backgroundImage!=='none'||['Top','Right','Bottom','Left'].some(s=>px(cs['border'+s+'Width'])>0&&cs['border'+s+'Style']!=='none'&&colorAlpha(cs['border'+s+'Color'])>.001)||el.tagName==='IMG'};
  const elements=[];let order=0,uid=0;
  const walk=el=>{
    if(!(el instanceof Element)||!visible(el)||el.namespaceURI==='http://www.w3.org/2000/svg')return;
    const cs=getComputedStyle(el),r=el.getBoundingClientRect();
    const btl=cs.borderTopLeftRadius;
    const it={id:idOf(el),order:dom.indexOf(el),zIndex:zi(el),stackPath:stackPath(el),tag:el.tagName.toLowerCase(),classes:[...el.classList],geometry:el.getAttribute('data-geometry')||'',rect:rect(r),opacity:opacityOf(el),rotation:rot(cs.transform),
      clipPath:cs.clipPath,backgroundColor:cs.backgroundColor,backgroundImage:cs.backgroundImage,backgroundUrl:bgUrl(cs.backgroundImage),borderRadii:[cs.borderTopLeftRadius,cs.borderTopRightRadius,cs.borderBottomRightRadius,cs.borderBottomLeftRadius],borderRadius:btl,borderRadiusPx:radiusPx(btl,r),boxShadow:cs.boxShadow,
      borderTop:{width:px(cs.borderTopWidth),color:cs.borderTopColor,style:cs.borderTopStyle},borderRight:{width:px(cs.borderRightWidth),color:cs.borderRightColor,style:cs.borderRightStyle},
      borderBottom:{width:px(cs.borderBottomWidth),color:cs.borderBottomColor,style:cs.borderBottomStyle},borderLeft:{width:px(cs.borderLeftWidth),color:cs.borderLeftColor,style:cs.borderLeftStyle},
      padding:pad(cs),fontFamily:cs.fontFamily,fontSize:px(cs.fontSize),fontWeight:cs.fontWeight,fontStyle:cs.fontStyle,color:cs.color,textDecoration:cs.textDecorationLine,lineHeightPx:lineHeightPx(cs),letterSpacingPx:px(cs.letterSpacing),
      textAlign:cs.textAlign,alignItems:cs.alignItems,justifyContent:cs.justifyContent,src:el.tagName==='IMG'?(el.currentSrc||el.src||''):null,innerText:''};
    const hasBorder=[it.borderTop,it.borderRight,it.borderBottom,it.borderLeft].some(b=>b.width>0&&b.style!=='none'&&colorAlpha(b.color)>0);
    const hasBg=colorAlpha(cs.backgroundColor)>.001&&cs.backgroundColor!=='transparent';
    const hasBgImg=cs.backgroundImage&&cs.backgroundImage!=='none';
    if(el===slide||hasBg||hasBgImg||hasBorder||cs.boxShadow!=='none'||el.tagName==='IMG'){
      elements.push(it);
    }
    [...el.children].forEach(walk);
  };walk(slide);

  const svgLines=[];

  const textRuns=[];
  const markers=[];
  return {deck:{w:sr.width,h:sr.height},elements,textRuns,markers,svgLines};
}
"""


def zkey(it):
    measured=it.get('domPaintOrder')
    if measured is not None:
        return ((int(measured['layer']),),int(measured['sequence']))
    path = it.get('stackPath')
    if not isinstance(path, list) or not path:
        path = [int(it.get('zIndex') or 0)]
    # Lexicographic stacking-context path followed by DOM/paint order.
    return (tuple(int(x or 0) for x in path), int(it.get('order') or 0))



def add_svg_shape(slide, item, sx, sy):
    """Supported SVG primitives remain individual editable native shapes."""
    r = item['rect']
    kind = MSO_SHAPE.OVAL if item['tag'] in ('circle','ellipse') else (MSO_SHAPE.ROUNDED_RECTANGLE if item.get('radius') else MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Inches(r['x']*sx), Inches(r['y']*sy), Inches(r['w']*sx), Inches(r['h']*sy))
    remove_theme_style(shape); clear_shape_effects(shape)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = min(.5, item['radius']/max(1e-9,min(r['w'],r['h'])))
    set_shape_fill(shape, item['fill'], item['opacity']*item['fillOpacity'])
    set_line(shape.line, item['stroke'], item['strokeWidth'], item['opacity']*item['strokeOpacity'])
    shape.line._get_or_add_ln().set('cap',{'round':'rnd','square':'sq'}.get(item.get('linecap'),'flat'))
    shape.name = 'kriss:' + item['id']
    return shape
